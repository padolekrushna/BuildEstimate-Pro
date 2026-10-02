from collections import OrderedDict
from io import BytesIO
from math import ceil
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AGENCY = "MAHARASHTRA INDUSTRIAL DEVELOPMENT CORPORATION"
AGENCY_SUBTITLE = "( A Government of Maharashtra Undertaking )"
HEADERS = ["Item No.", "Ref.to CSR", "Description of Item", "No.", "Length", "Breath", "Depth", "Quantity ", "Unit", "Rate", "Amount", "Openings Area", "Faces", "Coats", "Steps", "Riser", "Tread", "Waist Thick"]


def _description(item: Dict[str, Any]) -> str:
    return str(item.get("description") or item.get("subcategory") or "Work item").strip()


def _input_values(item: Dict[str, Any]) -> Dict[str, Any]:
    inputs = item.get("calculation", {}).get("inputs", {}) or {}
    count = inputs.get("component_count", inputs.get("count", inputs.get("no", 1)))
    opening_area = inputs.get("opening_area", inputs.get("openings_area", inputs.get("area", 0)))
    opening_line = item.get("work_category") == "Openings" or "opening" in str(item.get("subcategory", "")).lower()
    if opening_line and count:
        opening_area = float(opening_area or 0) / float(count)

    if "perimeter" in inputs:
        length = inputs.get("perimeter")
    elif "horizontal_run" in inputs:
        length = inputs.get("horizontal_run")
    elif "formwork_area" in inputs:
        length = inputs.get("formwork_area")
    elif "concrete_volume" in inputs:
        length = inputs.get("concrete_volume")
    elif "length" in inputs:
        length = inputs.get("length")
    elif "L" in inputs:
        length = inputs.get("L")
    else:
        length = 0

    breadth = inputs.get("width", inputs.get("thickness", inputs.get("W", inputs.get("steel_per_cum", 0))))
    depth = inputs.get("height", inputs.get("depth", inputs.get("thickness", inputs.get("H", inputs.get("waist_thickness", 0)))))
    return {
        "count": count,
        "length": length,
        "breadth": breadth,
        "depth": depth,
        "opening_area": opening_area,
        "faces": inputs.get("faces", 1),
        "coats": inputs.get("coats", 1),
        "steps": inputs.get("steps", 0),
        "riser": inputs.get("riser", 0),
        "tread": inputs.get("tread", 0),
        "waist": inputs.get("waist_thickness", 0),
    }


def _quantity_formula(item: Dict[str, Any], row: int) -> str:
    inputs = item.get("calculation", {}).get("inputs", {}) or {}
    category = item.get("work_category")
    subcategory = str(item.get("subcategory", "")).lower()

    if "horizontal_run" in inputs and "steps" in inputs and "waist_thickness" in inputs:
        return f"=D{row}*((F{row}*O{row}*P{row}*Q{row}/2)+(F{row}*SQRT(E{row}^2+(O{row}*P{row})^2)*R{row}))"
    if "steel_per_cum" in inputs and "concrete_volume" in inputs:
        return f"=E{row}*F{row}"
    if "formwork_area" in inputs:
        return f"=D{row}*E{row}"
    if category == "Openings" or "opening" in subcategory:
        return f"=D{row}*L{row}"
    if "perimeter" in inputs and "thickness" in inputs and "height" in inputs and "opening_volume" in inputs:
        return f"=D{row}*MAX(E{row}*F{row}*G{row}-L{row}*F{row},0)"
    if "perimeter" in inputs and "height" in inputs and ("opening_area" in inputs or "openings_area" in inputs):
        faces = "M" if "faces" in inputs else "1"
        coats = "N" if "coats" in inputs else "1"
        return f"=D{row}*MAX(E{row}*G{row}-L{row},0)*{faces}{row}*{coats}{row}"
    if "length" in inputs and "width" in inputs and "opening_area" in inputs and "faces" in inputs:
        return f"=D{row}*MAX(2*(E{row}+F{row})*G{row}-L{row},0)*M{row}*N{row}"
    if "length" in inputs and "width" in inputs and ("thickness" in inputs or "depth" in inputs or "height" in inputs):
        return f"=D{row}*E{row}*F{row}*G{row}"
    if "length" in inputs and "width" in inputs:
        return f"=D{row}*E{row}*F{row}"
    if "count" in inputs:
        return f"=D{row}"
    return f"=D{row}"


def _measurement_label(item: Dict[str, Any], index: int) -> str:
    inputs = item.get("calculation", {}).get("inputs", {}) or {}
    return str(inputs.get("location") or inputs.get("measurement_label") or inputs.get("component_name") or item.get("subcategory") or f"Measurement {index}")


def _group_items(items: List[Dict[str, Any]]):
    groups = OrderedDict()
    for item in items:
        key = (item.get("work_category"), item.get("subcategory"), item.get("dsr_item_no"), item.get("unit"), item.get("dsr_rate"))
        groups.setdefault(key, []).append(item)
    return list(groups.values())


def _format_sheet(ws, max_column: int):
    thin = Side(style="thin", color="777777")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=max_column):
        for cell in row:
            cell.font = Font(name="Bookman Old Style", size=11, color="202020")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4


def boq_to_excel_bytes(
    boq_items: List[Dict[str, Any]],
    project_summary: Dict[str, Any] = None,
    material_statement: List[Dict[str, Any]] = None,
    labour_statement: List[Dict[str, Any]] = None,
) -> bytes:
    """Create one formula-driven Class IV estimate with long item descriptions and measurement sections."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Class IV"
    summary = project_summary or {}
    ws.merge_cells("A1:R1")
    ws["A1"] = AGENCY
    ws.merge_cells("A2:R2")
    ws["A2"] = AGENCY_SUBTITLE
    ws["A1"].font = Font(name="Book Antiqua", size=12, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A2"].font = Font(name="Book Antiqua", size=12)
    ws["A2"].alignment = Alignment(horizontal="center")
    ws["A4"] = "Name of Work :-"
    ws["A4"].font = Font(name="Bookman Old Style", size=11, bold=True)
    ws["A4"].alignment = Alignment(horizontal="right")
    ws.merge_cells("C4:R4")
    ws["C4"] = summary.get("project_name", "")
    ws["C4"].font = Font(name="Bookman Old Style", size=11)
    ws.merge_cells("A5:R5")
    ws["A5"] = "Edit pale yellow measurement and rate cells. Click Quantity or Amount to inspect its formula; edits recalculate dependent totals in Excel."
    ws["A5"].font = Font(name="Bookman Old Style", size=9, italic=True, color="47534D")

    header_row = 7
    for column, header in enumerate(HEADERS, start=1):
        ws.cell(header_row, column, header)
    for cell in ws[header_row][:len(HEADERS)]:
        cell.font = Font(name="Bookman Old Style", size=12, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    say_rows = []
    item_number = 0
    row = header_row + 1
    for group in _group_items(boq_items):
        item_number += 1
        first = group[0]
        description = _description(first)
        ws.cell(row, 1, item_number)
        ws.cell(row, 2, first.get("dsr_item_no"))
        ws.cell(row, 3, description)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=18)
        ws.cell(row, 3).font = Font(name="Bookman Old Style", size=11, bold=True, color="9C0006" if first.get("dsr_approval_required") else "202020")
        ws.cell(row, 3).alignment = Alignment(wrap_text=True, vertical="top")
        description_lines = max(2, ceil(len(description) / 175))
        ws.row_dimensions[row].height = min(405, max(30, description_lines * 15 + 6))
        row += 1

        measurement_start = row
        for measurement_index, item in enumerate(group, start=1):
            inputs = _input_values(item)
            ws.cell(row, 3, _measurement_label(item, measurement_index))
            ws.cell(row, 3).font = Font(name="Bookman Old Style", size=10)
            for column, key in ((4, "count"), (5, "length"), (6, "breadth"), (7, "depth"), (12, "opening_area"), (13, "faces"), (14, "coats"), (15, "steps"), (16, "riser"), (17, "tread"), (18, "waist")):
                ws.cell(row, column, inputs[key])
                ws.cell(row, column).fill = PatternFill("solid", fgColor="FFF8E1")
            ws.cell(row, 8, _quantity_formula(item, row))
            ws.cell(row, 8).fill = PatternFill("solid", fgColor="E7EEE9")
            ws.cell(row, 8).number_format = "0.00"
            row += 1

        subtotal_row = row
        ws.cell(subtotal_row, 7, "Total")
        ws.cell(subtotal_row, 8, f"=SUM(H{measurement_start}:H{subtotal_row - 1})")
        ws.cell(subtotal_row, 8).font = Font(name="Bookman Old Style", size=11, bold=True)
        row += 1

        say_row = row
        precision = 0 if str(first.get("unit", "")).lower() in {"nos", "no", "m3", "cum"} else 1
        ws.cell(say_row, 7, "Say")
        ws.cell(say_row, 8, f"=ROUNDUP(H{subtotal_row},{precision})")
        ws.cell(say_row, 8).font = Font(name="Bookman Old Style", size=11, bold=True)
        ws.cell(say_row, 9, first.get("unit"))
        ws.cell(say_row, 10, first.get("dsr_rate"))
        ws.cell(say_row, 10).fill = PatternFill("solid", fgColor="FFF8E1")
        ws.cell(say_row, 11, f'=IF(OR(H{say_row}="",J{say_row}=""),"",H{say_row}*J{say_row})')
        ws.cell(say_row, 11).font = Font(name="Bookman Old Style", size=11, bold=True)
        ws.cell(say_row, 11).fill = PatternFill("solid", fgColor="E7EEE9")
        say_rows.append(say_row)
        row += 2

    total_row = row
    ws.cell(total_row, 8, "Total")
    ws.cell(total_row, 11, f"=SUM({','.join(f'K{say_row}' for say_row in say_rows)})" if say_rows else "=0")
    ws.cell(total_row, 8).font = Font(name="Bookman Old Style", size=11, bold=True)
    ws.cell(total_row, 11).font = Font(name="Bookman Old Style", size=11, bold=True)
    ws.cell(total_row, 11).number_format = '#,##0.00'

    widths = [8, 13, 43, 8, 10, 9, 10, 12, 9, 12, 14, 12, 8, 8, 8, 8, 8, 10]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    _format_sheet(ws, len(HEADERS))
    ws.row_dimensions[header_row].height = 34
    ws.freeze_panes = f"A{header_row + 1}"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
